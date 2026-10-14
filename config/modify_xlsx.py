from openpyxl import load_workbook

# 打开文件
# wb = load_workbook("/home/liyang/ETguider/Photosim6ft/et_100_det_inputs_1h_guider.xlsx")
wb = load_workbook("/home/liyang/cluster_mount/cluster_ET7/ET-fsgsim/config/et_inputs_transit_guider.xlsx")
# wb = load_workbook("/home/liyang/ETguider/Photosim6ft/et_inputs_microlens_guider.xlsx")
ws = wb.active
# print(ws["C23"].value)

# 修改单元格
# ws["C3"] = 35
# ws['C24'] = 0.2598 # pixel scale asec/pix for f=5160mm, pixelwidth=0.0065mm, microlens
# ws["C18"] = 0.676
# ws["C19"] = 0.127
# ws["C23"] = 1 
# ws["C27"] = 3
# ws["C29"] = 3
ws['C35'] = 100

# ws["C24"] = 3.146 # pixel scale asec/pix
# ws["C25"] = 6.5   # pixel width um
# ws["C41"] = 0.25  # exposure duration s
# # ws['C43'] = 1  # observing duration s
# ws['C15'] = 0.0
# ws['C16'] = 0.0
# ws['C18'] = 10.985
# ws['C19'] = 2.1
# ws['C30'] = 0.7
# ws['C33'] = 7
# ws['C34'] = False


# 保存
wb.save("/home/liyang/cluster_mount/cluster_ET7/ET-fsgsim/config/et_inputs_transit_guider.xlsx")
wb.close()

print("修改完成")